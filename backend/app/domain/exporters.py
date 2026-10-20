"""Report files: CSV (UTF-8 with a BOM), XLSX (openpyxl) and PDF (ReportLab, A4 landscape, DemandIQ colours)."""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime

CONTENT_TYPES = {
    "CSV": "text/csv;charset=utf-8",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "PDF": "application/pdf",
}
PDF_ROW_LIMIT = 2000


def file_slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def file_name(name: str, fmt: str) -> str:
    return f"{file_slug(name)}.{fmt.lower()}"


def build_csv(rows: list[dict]) -> bytes:
    headers = list(rows[0].keys())
    out = io.StringIO()
    writer = csv.writer(out, quoting=csv.QUOTE_ALL, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if row.get(h) is None else row.get(h) for h in headers])
    return ("\ufeff" + out.getvalue().rstrip("\n")).encode("utf-8")


def build_xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = list(rows[0].keys())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F7A8C")
    for row in rows:
        sheet.append(["" if row.get(h) is None else row.get(h) for h in headers])
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(row.get(header) or "")) for row in rows[:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(60, width + 2)
    sheet.freeze_panes = "A2"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def build_pdf(rows: list[dict], name: str, metadata: list[str], generated: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    jet, teal, lavender = colors.HexColor("#022b3a"), colors.HexColor("#1f7a8c"), colors.HexColor("#f0f2f8")
    headers = list(rows[0].keys())
    out = io.BytesIO()
    page_w, page_h = landscape(A4)

    def chrome(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(jet)
        canvas.rect(0, page_h - 58, page_w, 58, stroke=0, fill=1)
        canvas.setFillColor(teal)
        canvas.rect(0, page_h - 61, page_w, 3, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawString(36, page_h - 30, name)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(36, page_h - 46, "DemandIQ · Demand Forecasting & Inventory Management")
        canvas.setFillColor(colors.HexColor("#4d7585"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(36, 16, f"Generated {generated}")
        canvas.drawRightString(page_w - 36, 16, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=36, rightMargin=36, topMargin=72, bottomMargin=34, title=name, author="DemandIQ")
    meta_style = ParagraphStyle("meta", fontName="Helvetica", fontSize=8, leading=11, textColor=jet)
    cell_style = ParagraphStyle("cell", fontName="Helvetica", fontSize=7, leading=8.6, textColor=jet)
    head_style = ParagraphStyle("head", parent=cell_style, fontName="Helvetica-Bold", textColor=colors.white)
    story = [Paragraph(line, meta_style) for line in metadata] + [Spacer(1, 8)]
    body = [[Paragraph(str(h), head_style) for h in headers]]
    for row in rows[:PDF_ROW_LIMIT]:
        body.append([Paragraph(_escape("" if row.get(h) is None else str(row.get(h))), cell_style) for h in headers])
    table = LongTable(body, repeatRows=1, colWidths=[(page_w - 72) / len(headers)] * len(headers))
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), teal),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, lavender]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)
    if len(rows) > PDF_ROW_LIMIT:
        story += [Spacer(1, 10), Paragraph(f"Showing {PDF_ROW_LIMIT:,} of {len(rows):,} records. Use XLSX or CSV for the full dataset.", meta_style)]
    doc.build(story, onFirstPage=chrome, onLaterPages=chrome)
    return out.getvalue()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_report_file(fmt: str, rows: list[dict], name: str, metadata: list[str], generated: str | None = None) -> tuple[bytes, str, str]:
    """→ (content, file name, content type). Raises ValueError when there are no rows."""
    if not rows:
        raise ValueError("There is no matching data to export.")
    if fmt == "CSV":
        content = build_csv(rows)
    elif fmt == "XLSX":
        content = build_xlsx(rows)
    else:
        content = build_pdf(rows, name, metadata, generated or datetime.now().strftime("%d %b %Y, %H:%M"))
    return content, file_name(name, fmt), CONTENT_TYPES[fmt]
