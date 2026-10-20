"""Reads an uploaded sales file of any layout into daily sales records: one record per product per day.

A port of importSales (src/lib/salesImport.js) and of the file reading in src/services/mock/routes/sales.js,
so FastAPI accepts exactly the files the demo backend accepts (docs/API_CONTRACT.md, section 7):
columns found by their names in any order, a header under titles, no header, dates across the top, dates
and numbers in the usual forms, rows for the same product and day added up, new products collected.
"""
from __future__ import annotations

import calendar
import io
import math
import re
import time
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta

MAX_NEW_PRODUCTS = 10000


def _nfkd(text: str) -> str:
    return unicodedata.normalize("NFKD", text)


def header_key(text) -> str:
    """"Units Sold" → "unitssold" """
    value = _nfkd(str(text if text is not None else "").lstrip("\ufeff")).lower()
    return re.sub(r"[^a-z0-9%]", "", value)


def product_key(text) -> str:
    """"pp 4012" and "PP-4012" are the same product."""
    return re.sub(r"[^A-Z0-9]", "", _nfkd(str(text if text is not None else "")).upper())


# ------------------------------------------------------------------------------------------------
# Columns
# ------------------------------------------------------------------------------------------------


def _has(pattern: str, k: str) -> bool:
    return re.search(pattern, k) is not None


ROLES = [
    (
        "date",
        ["date", "salesdate", "saledate", "orderdate", "transactiondate", "txndate", "invoicedate", "billdate", "day", "datetime", "timestamp", "period", "documentdate", "postingdate", "dateofsale", "datum", "fecha"],
        lambda k: "date" in k and not _has(r"(ship|due|deliver|expir|birth|join|create|update|modif|start|end)", k),
    ),
    (
        "units",
        ["unitssold", "units", "quantity", "quantitysold", "qty", "qtysold", "soldqty", "soldquantity", "salesqty", "salesquantity", "salesunits", "unitsales", "unitsold", "noofunits", "numberofunits", "quantityordered", "orderquantity", "orderqty", "pieces", "pcs", "sold", "itemssold", "itemsold", "volume", "menge", "anzahl", "cantidad", "quantite", "qte", "quantita"],
        lambda k: _has(r"(qty|quantity|unitssold|unitsold|unitsales|sold)", k) and not _has(r"(forecast|stock|inventory|return|price|cost|purchase|receiv|onhand|avail|reorder)", k),
    ),
    ("discountAmount", ["discountamount", "discountamt", "discountvalue"], lambda k: _has("disc", k) and _has(r"(amount|amt|value)", k)),
    ("discount", ["discountrate", "discount", "discountpercent", "discountpercentage", "discountpct", "disc", "discount%", "disc%", "discountrate%"], lambda k: k.startswith("disc")),
    (
        "price",
        ["unitprice", "price", "sellingprice", "saleprice", "salesprice", "priceperunit", "unitsellingprice", "itemprice", "priceeach", "rate", "unitrate", "mrp", "listprice", "preis", "precio", "prix", "prezzo"],
        lambda k: "price" in k and not _has(r"(competitor|purchase|cost|total)", k),
    ),
    (
        "revenue",
        ["sales", "revenue", "netrevenue", "netsales", "salesamount", "salesvalue", "amount", "totalamount", "totalsales", "totalrevenue", "netamount", "lineamount", "linetotal", "total", "turnover", "value", "umsatz", "betrag", "ventas", "importe", "ventes", "montant"],
        lambda k: _has(r"(revenue|sales|amount|turnover)", k) and not _has(r"(qty|quantity|unit|count|tax|cost|profit|margin|target|forecast|person|rep|channel)", k),
    ),
    (
        "sku",
        ["sku", "productid", "skuid", "skucode", "productcode", "productsku", "itemid", "itemcode", "itemno", "itemnumber", "productno", "productnumber", "partno", "partnumber", "materialcode", "materialno", "articlecode", "articleno", "articlenumber", "stockcode", "upc", "ean", "barcode", "asin", "code", "artikelnummer", "artikelnr", "codigo"],
        lambda k: _has(r"(sku|barcode)", k) or (_has(r"(product|item|article|material|part)", k) and _has(r"(id|code|no|num|nbr|number|key)$", k)),
    ),
    (
        "name",
        ["productname", "itemname", "productdescription", "itemdescription", "description", "productdesc", "itemdesc", "producttitle", "particulars", "name", "title"],
        lambda k: _has(r"(product|item|article)", k) and _has(r"(name|desc|title)", k),
    ),
    (
        "product",
        ["product", "item", "products", "items", "article", "material", "artikel", "produkt", "producto", "produit", "articulo", "prodotto"],
        lambda k: _has(r"^(product|item)", k) and not _has(r"(categor|type|group|class|line|family)", k),
    ),
    (
        "stock",
        ["stock", "stocklevel", "stockonhand", "onhand", "onhandqty", "inventory", "inventorylevel", "inventoryonhand", "closingstock", "currentstock", "availablestock", "stockqty", "soh"],
        lambda k: _has(r"(stock|inventory|onhand)", k) and not _has(r"(order|forecast|out|value|cost)", k),
    ),
    (
        "category",
        ["category", "productcategory", "categoryname", "itemcategory", "productgroup", "productfamily", "productline", "producttype", "itemtype", "department", "dept", "type"],
        lambda k: "categor" in k,
    ),
    ("region", ["region", "salesregion", "area", "zone", "territory", "state", "province", "market", "country"], lambda k: _has(r"(region|zone|territor|province)", k)),
    (
        "location",
        ["location", "store", "storeid", "storename", "storecode", "storeno", "warehouse", "warehouseid", "branch", "branchname", "outlet", "outletname", "site", "shop", "shopname", "depot", "plant", "locationid", "city"],
        lambda k: _has(r"(store|warehouse|branch|outlet|location|depot)", k),
    ),
    ("channel", ["channel", "saleschannel", "orderchannel", "distributionchannel"], lambda k: "channel" in k),
]


def roles_from_header(header: list) -> dict[str, int]:
    keys = [header_key(cell) for cell in header]
    found: dict[str, int] = {}
    taken: set[int] = set()
    for role, exact, _ in ROLES:
        for name in exact:
            index = next((i for i, k in enumerate(keys) if k == name and i not in taken), -1)
            if index >= 0:
                found[role] = index
                taken.add(index)
                break
    for role, _, partial in ROLES:
        if role in found:
            continue
        index = next((i for i, k in enumerate(keys) if k and i not in taken and partial(k)), -1)
        if index >= 0:
            found[role] = index
            taken.add(index)
    return found


def _product_roles(found: dict) -> int:
    return sum(1 for role in ("sku", "product", "name") if role in found)


def _cell(row: list, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else str(value).strip()


def _looks_like_header(rows: list[list]) -> bool:
    def values(row):
        return [str(c if c is not None else "").strip() for c in row if str(c if c is not None else "").strip()]

    def is_data(value):
        return read_date(value) is not None or read_number(value) is not None

    return len(rows) > 1 and len(values(rows[0])) > 0 and not any(is_data(v) for v in values(rows[0])) and any(is_data(v) for v in values(rows[1]))


def _find_header(rows: list[list]):
    best = None
    for i in range(min(len(rows), 15)):
        found = roles_from_header(rows[i])
        score = int("date" in found) + int(_product_roles(found) > 0) + int("units" in found)
        if score >= 2 and (best is None or score > best["score"]):
            best = {"index": i, "found": found, "score": score}
    return best


def _find_wide_header(rows: list[list]):
    for i in range(min(len(rows) - 1, 15)):
        cells = [str(c if c is not None else "").strip() for c in rows[i]]
        order = date_order(cells)
        dates = []
        for index, text in enumerate(cells):
            parsed = read_date(text, order)
            if parsed:
                dates.append({"index": index, "date": parsed, "text": text})
        if len(dates) < 2:
            continue
        is_date = {d["index"] for d in dates}
        labels = [t for idx, t in enumerate(cells) if t and idx not in is_date]
        roles = roles_from_header(["" if idx in is_date else t for idx, t in enumerate(cells)])
        if any(read_number(t) is not None for t in labels) or "date" in roles or "units" in roles:
            continue
        for row in rows[i + 1 : i + 6]:
            if any(read_number(row[d["index"]] if d["index"] < len(row) else "") is not None for d in dates):
                return {"index": i, "dates": dates}
    return None


def _unpivot(rows: list[list], wide: dict) -> list[list]:
    head = [str(c if c is not None else "").strip() for c in rows[wide["index"]]]
    is_date = {d["index"] for d in wide["dates"]}
    width = max(len(row) for row in rows)
    keep = [i for i in range(width) if i not in is_date]
    table = [[(head[i] if i < len(head) and head[i] else f"Column {i + 1}") for i in keep] + ["Date", "Units"]]
    for row in rows[wide["index"] + 1 :]:
        for d in wide["dates"]:
            value = _cell(row, d["index"])
            if value:
                table.append([(row[i] if i < len(row) and row[i] is not None else "") for i in keep] + [d["date"], value])
    return table


# ------------------------------------------------------------------------------------------------
# Values
# ------------------------------------------------------------------------------------------------

MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12}


def _month(word: str):
    w = word.lower()
    return MONTHS.get(w[:4] if w.startswith("sept") else w[:3])


def _full_year(y: str) -> int:
    return (2000 if int(y) < 70 else 1900) + int(y) if len(y) == 2 else int(y)


def _iso_of(y: int, m: int, d: int) -> str | None:
    if not (1950 <= y <= 2100 and 1 <= m <= 12 and 1 <= d <= 31):
        return None
    if d > calendar.monthrange(y, m)[1]:
        return None
    return f"{y:04d}-{m:02d}-{d:02d}"


NUMERIC_DATE = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})(?:[ T,].*)?$")


def date_order(values) -> str:
    day_first = month_first = 0
    for value in values:
        m = NUMERIC_DATE.match(str(value).strip())
        if not m:
            continue
        if int(m.group(1)) > 12:
            day_first += 1
        if int(m.group(2)) > 12:
            month_first += 1
    return "MDY" if month_first > day_first else "DMY"


def excel_serial_date(value) -> str | None:
    """An Excel day number (45834) → "2025-06-26" (1900 date system)."""
    try:
        serial = float(value)
    except (TypeError, ValueError):
        return None
    days = math.floor(serial)
    return (date(1899, 12, 30) + timedelta(days=days)).isoformat()


def read_date(value, order: str = "DMY") -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value if value is not None else "").strip()
    if not text:
        return None
    text = re.sub(r"^(mon|tue|wed|thu|fri|sat|sun)[a-z]*\.?,?\s+", "", text, flags=re.I)
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T].*)?$", text)
    if m:
        return _iso_of(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = NUMERIC_DATE.match(text)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), _full_year(m.group(3))
        return _iso_of(y, a, b) if order == "MDY" else _iso_of(y, b, a)
    m = re.match(r"^(\d{1,2})(?:st|nd|rd|th)?[-\s./]*([A-Za-z]{3,9})\.?[-\s.,/]*(\d{2}|\d{4})(?:[ T,].*)?$", text)
    if m and _month(m.group(2)):
        return _iso_of(_full_year(m.group(3)), _month(m.group(2)), int(m.group(1)))
    m = re.match(r"^([A-Za-z]{3,9})\.?[-\s.]*(\d{1,2})(?:st|nd|rd|th)?,?[-\s./]*(\d{2}|\d{4})(?:[ T,].*)?$", text)
    if m and _month(m.group(1)):
        return _iso_of(_full_year(m.group(3)), _month(m.group(1)), int(m.group(2)))
    m = re.match(r"^(\d{4})(\d{2})(\d{2})$", text)
    if m:
        return _iso_of(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    if re.match(r"^\d{5}(\.\d+)?$", text) and 18000 <= float(text) <= 73000:
        return excel_serial_date(text)
    return None


CURRENCY = re.compile(r"^(rs\.?|inr|usd|sar|aed|eur|gbp|qar|kwd|omr|bhd)\s*|\s*(rs\.?|inr|usd|sar|aed|eur|gbp|qar|kwd|omr|bhd)$", re.I)


def read_number(value) -> dict | None:
    """→ {"value": float, "percent": bool} or None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return {"value": float(value), "percent": False} if math.isfinite(value) else None
    text = str(value if value is not None else "").strip()
    if not text:
        return None
    negative = False
    if re.match(r"^\(.*\)$", text):
        negative = True
        text = text[1:-1].strip()
    percent = re.search(r"%\s*$", text) is not None
    text = re.sub(r"%\s*$", "", text)
    text = CURRENCY.sub("", text)
    text = re.sub(r"[₹$€£¥\s\u00a0']", "", text)
    if re.match(r"^[-+]", text):
        negative = negative != text.startswith("-")
        text = text[1:]
    if re.match(r"^(\d+\.?\d*|\.\d+)e[-+]?\d+$", text, re.I):
        number = float(text) * (-1 if negative else 1)
        return {"value": number, "percent": percent} if math.isfinite(number) else None
    if not re.match(r"^[\d.,]+$", text):
        return None
    if re.match(r"^\d+,\d{1,2}$", text):
        text = text.replace(",", ".", 1)
    elif re.match(r"^\d{1,3}(\.\d{3})+(,\d+)?$", text) and "," in text:
        text = text.replace(".", "").replace(",", ".", 1)
    else:
        text = text.replace(",", "")
    if not re.match(r"^(\d+\.?\d*|\.\d+)$", text):
        return None
    number = float(text) * (-1 if negative else 1)
    return {"value": number, "percent": percent} if math.isfinite(number) else None


def _round(value: float, digits: int) -> float:
    """Number(value.toFixed(digits))"""
    return float(f"{value:.{digits}f}")


def _median(values: list[float]):
    if not values:
        return None
    ordered = sorted(values)
    middle = len(ordered) // 2
    return ordered[middle] if len(ordered) % 2 else (ordered[middle - 1] + ordered[middle]) / 2


def _most_common(counts: Counter):
    best = None
    for value, n in counts.items():  # insertion order: the first one wins a tie
        if best is None or n > counts[best]:
            best = value
    return best


ID_LIKE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._\-/]{0,39}$")


def _new_sku(id_text: str, name: str, taken: set[str]) -> str:
    as_id = id_text.upper() if id_text and ID_LIKE.match(id_text) else None
    base = as_id
    if base is None:
        base = re.sub(r"[^A-Z0-9]+", "-", _nfkd(str(name or id_text or "PRODUCT")).upper()).strip("-")[:30] or "PRODUCT"
    sku = base
    n = 2
    while product_key(sku) in taken:
        sku = f"{base}-{n}"
        n += 1
    taken.add(product_key(sku))
    return sku


def _listed(items: list[str]) -> str:
    return f"{', '.join(items[:-1])} and {items[-1]}" if len(items) > 1 else items[0]


COLUMN_HINTS = {
    "date": "a date column (Date, Order Date …)",
    "product": "a product column (SKU, Product ID, Item or Product Name …)",
    "units": "a quantity column (Units Sold, Quantity, Qty …)",
}

SUMMARY = re.compile(r"^(grand\s*|sub\s*-?\s*)?totals?\s*:?$", re.I)


def _share_of(body: list[list], index: int | None, test) -> float:
    values = [v for v in (_cell(row, index) for row in body[:500]) if v]
    return sum(1 for v in values if test(v)) / len(values) if values else 0.0


def _is_whole(text: str) -> bool:
    n = read_number(text)
    return bool(n) and not n["percent"] and n["value"] >= 0 and float(_round(n["value"], 6)).is_integer()


def _is_text(text: str) -> bool:
    return read_number(text) is None and read_date(text) is None


def _is_date(text: str) -> bool:
    return read_date(text, "DMY") is not None or read_date(text, "MDY") is not None


def _column_of_dates(body: list[list], columns: list[int]):
    for i in columns:
        if _share_of(body, i, lambda v: not re.match(r"^\d+(\.\d+)?$", v) and _is_date(v)) >= 0.8:
            return i
    for i in columns:
        if _share_of(body, i, _is_date) >= 0.8:
            return i
    return None


def _empty(errors: list[str], total: int = 0, columns: dict | None = None) -> dict:
    return {
        "total": total, "valid": 0, "combined": 0, "missing": 0, "duplicates": 0, "invalidDates": 0, "invalidNumeric": 0,
        "errors": errors, "records": [], "newProducts": [], "columns": columns or {},
    }


def _js_int(value: float) -> int:
    return int(math.floor(value + 0.5))


def _num(value: float) -> float | int:
    """Whole numbers as int (like JSON numbers in JS)."""
    return int(value) if float(value).is_integer() else value


def import_sales(rows: list[list], catalog: list[dict]) -> dict:
    """Validates the rows of a sales file against the catalog ({sku, name, category, location, region,
    channel, unitPrice}). → {total, valid, combined, missing, duplicates, invalidDates, invalidNumeric,
    errors, records, newProducts, columns}"""
    filled = [row for row in rows if any(str(c if c is not None else "").strip() != "" for c in row)]
    if not filled:
        return _empty(["The file is empty."])

    header = _find_header(filled)
    across = None
    if not header:
        wide = _find_wide_header(filled)
        if wide:
            def label(d):
                return d["date"] if re.match(r"^\d+(\.\d+)?$", d["text"]) else d["text"]

            across = f"{label(wide['dates'][0])} … {label(wide['dates'][-1])}"
            filled = _unpivot(filled, wide)
            found_wide = roles_from_header(filled[0])
            found_wide.pop("revenue", None)
            found_wide.pop("discountAmount", None)
            header = {"index": 0, "found": found_wide}
    if not header and _looks_like_header(filled):
        header = {"index": 0, "found": roles_from_header(filled[0])}
    if header:
        found = dict(header["found"])
        names = [(str(c if c is not None else "").lstrip("\ufeff").strip() or f"Column {i + 1}") for i, c in enumerate(filled[header["index"]])]
        body = filled[header["index"] + 1 :]
    else:
        body = filled
        width0 = max(len(row) for row in body)
        names = [f"Column {i + 1}" for i in range(width0)]
        found = _roles_from_contents(body, width0, catalog)
    width = max([len(names)] + [len(row) for row in body[:500]])

    def free():
        used = set(found.values())
        return [i for i in range(width) if i not in used]

    if "date" not in found:
        dates = _column_of_dates(body, free())
        if dates is not None:
            found["date"] = dates
    if (
        "units" not in found
        and "price" not in found
        and "revenue" in found
        and re.search(r"(sales|sold)", header_key(names[found["revenue"]] if found["revenue"] < len(names) else ""))
        and _share_of(body, found["revenue"], _is_whole) >= 0.98
    ):
        found["units"] = found.pop("revenue")

    known = set()
    for p in catalog:
        known.add(product_key(p["sku"]))
        known.add(product_key(p["name"]))
    identifier = found.get("sku", found.get("product", found.get("name")))
    if identifier is None:
        identifier = next((i for i in free() if _share_of(body, i, lambda v: product_key(v) in known) >= 0.5), None)
    if identifier is None and across:
        identifier = next((i for i in free() if _share_of(body, i, _is_text) >= 0.8), None)
    if identifier is None:
        identifier = found.get("category")
    name_column = next((index for index in (found.get("name"), found.get("product")) if index is not None and index != identifier), None)

    merged = {**found, "product": identifier, "name": name_column}
    merged.pop("sku", None)
    columns = {role: names[index] for role, index in merged.items() if index is not None and index < len(names)}
    if across:
        columns["date"] = across
        columns["units"] = across
    lacking = [role for role, missing in (("date", "date" not in found), ("product", identifier is None), ("units", "units" not in found)) if missing]
    if lacking:
        more = " …" if len(names) > 12 else ""
        return _empty(
            [f"Couldn't find {_listed([COLUMN_HINTS[r] for r in lacking])}. Columns in the file: {', '.join(names[:12])}{more}."],
            len(body),
            columns,
        )
    body = [row for row in body if not any(SUMMARY.match(str(c if c is not None else "").strip()) for c in row)]
    if not body:
        return _empty(["The file has a header but no rows."], 0, columns)

    order = date_order([_cell(row, found["date"]) for row in body[:2000]])
    discount_percent = "discount" in found and any(((read_number(_cell(row, found["discount"])) or {}).get("value", 0)) > 1 for row in body)

    by_sku = {product_key(p["sku"]): p for p in catalog}
    by_name = {product_key(p["name"]): p for p in catalog}
    taken = set(by_sku.keys())
    new_by: dict[str, dict] = {}
    groups: dict[str, dict] = {}
    seen: set[str] = set()
    counts = {"missing": 0, "duplicates": 0, "invalidDates": 0, "invalidNumeric": 0, "combined": 0}
    examples: dict[str, object] = {}

    def note(kind, value):
        counts[kind] += 1
        examples.setdefault(kind, value)

    for row_index, row in enumerate(body):
        signature = "\u0001".join(str(c if c is not None else "").strip() for c in row)
        if signature in seen:
            note("duplicates", row_index)
            continue
        seen.add(signature)
        date_text = _cell(row, found["date"])
        id_text = _cell(row, identifier)
        name = _cell(row, name_column)
        units_text = _cell(row, found["units"])
        if not date_text or not (id_text or name) or units_text == "":
            note("missing", row_index)
            continue
        sale_date = read_date(date_text, order)
        if not sale_date:
            note("invalidDates", date_text)
            continue
        units = read_number(units_text)
        if not units or units["value"] < 0 or not float(_round(units["value"], 6)).is_integer():
            note("invalidNumeric", units_text)
            continue
        u = _js_int(units["value"])

        product = by_sku.get(product_key(id_text)) or by_name.get(product_key(id_text))
        if not product and name:
            product = by_name.get(product_key(name)) or by_sku.get(product_key(name))

        price_text = _cell(row, found.get("price"))
        discount_text = _cell(row, found.get("discount"))
        amount_text = _cell(row, found.get("discountAmount"))
        revenue_text = _cell(row, found.get("revenue"))
        price = read_number(price_text) if price_text else None
        if price_text and (not price or price["value"] < 0):
            note("invalidNumeric", price_text)
            continue
        discount = 0.0
        if discount_text:
            dn = read_number(discount_text)
            if not dn:
                note("invalidNumeric", discount_text)
                continue
            discount = dn["value"] / 100 if dn["percent"] or discount_percent else dn["value"]
            if discount < 0 or discount > 1:
                note("invalidNumeric", discount_text)
                continue
        p = price["value"] if price else None
        if p is None and revenue_text:
            revenue = read_number(revenue_text)
            if not revenue or revenue["value"] < 0:
                note("invalidNumeric", revenue_text)
                continue
            if u > 0:
                p = revenue["value"] / u / (1 - min(discount, 0.9999))
        if p is None:
            p = float(product["unitPrice"]) if product and product.get("unitPrice") is not None else 0.0
        if amount_text and u > 0 and p > 0:
            amount = read_number(amount_text)
            if not amount or amount["value"] < 0 or amount["value"] > u * p:
                note("invalidNumeric", amount_text)
                continue
            discount = amount["value"] / (u * p)

        fresh = None
        if not product:
            k = product_key(id_text or name)
            fresh = new_by.get(k)
            if not fresh:
                fresh = {
                    "key": k,
                    "id": id_text,
                    "name": name or ("" if ID_LIKE.match(id_text or "") else id_text),
                    "attributes": {"category": Counter(), "region": Counter(), "location": Counter(), "channel": Counter()},
                    "prices": [],
                    "stock": {},
                }
                new_by[k] = fresh
            for role in ("category", "region", "location", "channel"):
                value = _cell(row, found.get(role))
                if value:
                    fresh["attributes"][role][value] += 1
            if p > 0:
                fresh["prices"].append(p)
            stock = read_number(_cell(row, found.get("stock")))
            if stock and stock["value"] >= 0:
                fresh["stock"][sale_date] = fresh["stock"].get(sale_date, 0) + _js_int(stock["value"])

        group_key = f"{sale_date}|{product['sku'] if product else 'new:' + fresh['key']}"
        group = groups.get(group_key)
        gross = u * p
        if group:
            counts["combined"] += 1
            group["units"] += u
            group["gross"] += gross
            group["net"] += gross * (1 - discount)
        else:
            groups[group_key] = {"date": sale_date, "product": product, "fresh": fresh, "units": u, "gross": gross, "net": gross * (1 - discount), "first": row_index}

    if len(new_by) > MAX_NEW_PRODUCTS:
        return _empty(
            [f"The file names {len(new_by):,} products that aren't in the catalog (at most {MAX_NEW_PRODUCTS:,} at a time). Check that \"{columns.get('product')}\" is the product column."],
            len(body),
            columns,
        )
    new_products = []
    for fresh in new_by.values():
        stock_dates = sorted(fresh["stock"].keys())
        product = {
            "sku": _new_sku(fresh["id"], fresh["name"], taken),
            "name": fresh["name"] or fresh["id"],
            "category": _most_common(fresh["attributes"]["category"]) or "Uncategorized",
            "region": _most_common(fresh["attributes"]["region"]) or "Unassigned",
            "location": _most_common(fresh["attributes"]["location"]) or "Unassigned",
            "channel": _most_common(fresh["attributes"]["channel"]) or "Unassigned",
            "unitPrice": _num(_round(_median(fresh["prices"]) or 0, 2)),
            "stock": fresh["stock"][stock_dates[-1]] if stock_dates else None,
        }
        fresh["product"] = product
        new_products.append(product)

    stamp = int(time.time() * 1000)
    records = []
    for group in groups.values():
        product = group["product"] or group["fresh"]["product"]
        unit_price = _round(group["gross"] / group["units"], 2) if group["units"] else _round(float(product.get("unitPrice") or 0), 2)
        discount_rate = _round(1 - group["net"] / group["gross"], 4) if group["gross"] > 0 else 0
        records.append(
            {
                "id": f"UPL-{stamp}-{group['first']}",
                "date": group["date"],
                "sku": product["sku"],
                "category": product["category"],
                "location": product["location"],
                "region": product["region"],
                "channel": product["channel"],
                "units": group["units"],
                "unitPrice": _num(unit_price),
                "discountRate": _num(discount_rate),
                "netRevenue": _num(_round(group["units"] * unit_price * (1 - discount_rate), 2)),
            }
        )
    records.sort(key=lambda r: (r["date"], r["sku"]))
    result = {"total": len(body), "valid": len(records), **counts, "errors": [], "records": records, "newProducts": new_products if records else [], "columns": columns}
    if not records:
        result["errors"].append(_no_rows_message(len(body), counts, examples))
    return result


def _no_rows_message(total: int, counts: dict, examples: dict) -> str:
    def has(n):
        return f"{n:,} {'has' if n == 1 else 'have'}"

    texts = {
        "invalidDates": lambda n, e: f'{has(n)} a date that can\'t be read (such as "{e}")',
        "invalidNumeric": lambda n, e: f'{has(n)} a number that can\'t be used (such as "{e}")',
        "missing": lambda n, e: f"{has(n)} no date, product or quantity",
        "duplicates": lambda n, e: f"{n:,} {'repeats' if n == 1 else 'repeat'} an earlier row",
    }
    kinds = [k for k in ("invalidDates", "invalidNumeric", "missing", "duplicates") if counts[k]]
    kinds.sort(key=lambda k: -counts[k])  # stable: ties keep the order above
    return f"None of the {total:,} rows could be used: {'; '.join(texts[k](counts[k], examples.get(k)) for k in kinds)}."


def _roles_from_contents(body: list[list], width: int, catalog: list[dict]) -> dict[str, int]:
    known = set()
    for p in catalog:
        known.add(product_key(p["sku"]))
        known.add(product_key(p["name"]))
    found: dict[str, int] = {}

    def free():
        used = set(found.values())
        return [i for i in range(width) if i not in used]

    column = _column_of_dates(body, free())
    if column is not None:
        found["date"] = column
    product = next((i for i in free() if _share_of(body, i, lambda v: product_key(v) in known) >= 0.5), None)
    if product is None:
        for i in free():
            values = [v for v in (_cell(row, i) for row in body[:500]) if v]
            if _share_of(body, i, _is_text) >= 0.8 and len(set(values)) <= max(1, len(values) * 0.5):
                product = i
                break
    if product is not None:
        found["product"] = product
    units = next((i for i in free() if _share_of(body, i, _is_whole) >= 0.9), None)
    if units is not None:
        found["units"] = units
    return found


# ------------------------------------------------------------------------------------------------
# Files: what a file is (from its first bytes), text decoding, delimiters, Excel workbooks
# ------------------------------------------------------------------------------------------------

UNREADABLE = {
    "xls": "Excel 97–2003 (.xls) files can't be read. Save it as .xlsx or CSV and upload it again.",
    "pdf": "A PDF can't be read as sales data. Upload a CSV, text or Excel (.xlsx) file.",
    "binary": "Unsupported file type. Upload a CSV, text or Excel (.xlsx) file.",
}


def kind_of(data: bytes) -> str:
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:4] == b"\xd0\xcf\x11\xe0":
        return "xls"
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "text"
    return "binary" if b"\x00" in data[:4096] else "text"


def decode_text(data: bytes) -> str:
    if data[:2] == b"\xff\xfe":
        return data[2:].decode("utf-16-le", errors="replace")
    if data[:2] == b"\xfe\xff":
        return data[2:].decode("utf-16-be", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def split_delimited(text: str, delimiter: str = ",") -> list[list[str]]:
    rows: list[list[str]] = []
    row: list[str] = []
    cell: list[str] = []
    quoted = False
    body = text.lstrip("\ufeff")
    i = 0
    n = len(body)
    while i < n:
        ch = body[i]
        if quoted:
            if ch == '"' and i + 1 < n and body[i + 1] == '"':
                cell.append('"')
                i += 1
            elif ch == '"':
                quoted = False
            else:
                cell.append(ch)
        elif ch == '"':
            quoted = True
        elif ch == delimiter:
            row.append("".join(cell))
            cell = []
        elif ch in "\n\r":
            if ch == "\r" and i + 1 < n and body[i + 1] == "\n":
                i += 1
            row.append("".join(cell))
            rows.append(row)
            row = []
            cell = []
        else:
            cell.append(ch)
        i += 1
    if cell or row:
        row.append("".join(cell))
        rows.append(row)
    return [r for r in rows if any(c.strip() != "" for c in r)]


def detect_delimiter(text: str) -> str:
    lines = [line for line in re.split(r"\r?\n", text.lstrip("\ufeff")) if line.strip()][:25]
    best = {"delimiter": ",", "lines": 0, "cells": 1}
    for delimiter in [",", "\t", ";", "|"]:
        widths = []
        for line in lines:
            split = split_delimited(line, delimiter)
            width = len(split[0]) if split else 1
            if width > 1:
                widths.append(width)
        tally = Counter(widths)
        for cells, count in tally.items():
            if count > best["lines"] or (count == best["lines"] and cells > best["cells"]):
                best = {"delimiter": delimiter, "lines": count, "cells": cells}
    return best["delimiter"]


class WorkbookError(ValueError):
    pass


def read_xlsx_sheets(data: bytes) -> list[list[list]]:
    """Every worksheet (visible first, hidden last) as rows of cell texts; a date cell as its Excel day number."""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # not a readable workbook
        raise WorkbookError("This isn't a valid .xlsx file.") from error
    sheets = []
    ordered = sorted(workbook.worksheets, key=lambda ws: 0 if ws.sheet_state == "visible" else 1)
    epoch = datetime(1899, 12, 30)
    for sheet in ordered:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = []
            for value in row:
                if value is None:
                    cells.append("")
                elif isinstance(value, datetime):
                    delta = value - epoch
                    serial = delta.days + delta.seconds / 86400
                    cells.append(str(int(serial)) if float(serial).is_integer() else str(serial))
                elif isinstance(value, date):
                    cells.append(str((datetime(value.year, value.month, value.day) - epoch).days))
                elif isinstance(value, float) and value.is_integer():
                    cells.append(str(int(value)))
                else:
                    cells.append(str(value))
            rows.append(cells)
        sheets.append(rows)
    workbook.close()
    return sheets


def read_upload(data: bytes, catalog: list[dict]) -> dict:
    """Reads a file whatever its name: delimited text, or the first sheet of a workbook with sales on it."""
    if kind_of(data) != "xlsx":
        text = decode_text(data)
        return import_sales(split_delimited(text, detect_delimiter(text)), catalog)
    fullest = None
    for rows in read_xlsx_sheets(data):
        result = import_sales(rows, catalog)
        if result["valid"]:
            return result
        if fullest is None or result["total"] > fullest["total"]:
            fullest = result
    return fullest or _empty(["The workbook has no sheets."])


def validation_label(v: dict) -> str:
    if v["errors"] or v["valid"] == 0:
        return "Failed"
    warnings = v["missing"] + v["duplicates"] + v["invalidDates"] + v["invalidNumeric"]
    return f"{warnings} warning{'' if warnings == 1 else 's'}" if warnings else "Valid"
